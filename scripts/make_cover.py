"""Create the assignment cover sheet after the GitHub URL is available."""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def hyperlink(paragraph, label, url):
    relation = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "154574")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--repository-url", required=True)
    parser.add_argument("--output", type=Path, default=Path("submission/KrishnaTandon_230911516_FGAI_Assignment3.docx"))
    args = parser.parse_args()
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1.05)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.1)
    section.right_margin = Inches(1.1)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    title = document.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    for border in title._element.pPr.findall(qn("w:pBdr")):
        title._element.pPr.remove(border)

    heading = document.add_paragraph(style="Title")
    heading.add_run("Foundations of Generative AI Assignment 3")
    for border in heading._p.pPr.findall(qn("w:pBdr")) if heading._p.pPr is not None else []:
        heading._p.pPr.remove(border)
    summary = document.add_paragraph()
    summary.add_run("English to Hindi Machine Translation Comparison").bold = True
    summary.paragraph_format.space_after = Pt(22)

    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    entries = [
        ("Student name", "Krishna Tandon"),
        ("Registration number", "230911516"),
        ("Assignment title", "Foundations of Generative AI Assignment 3"),
        ("Developed system", "English to Hindi Machine Translation Comparison"),
        ("Dataset", "IIT Bombay English Hindi Parallel Corpus"),
        ("Dataset source", "https://huggingface.co/datasets/cfilt/iitb-english-hindi"),
        ("GitHub repository", args.repository_url),
    ]
    for label, value in entries:
        left, right = table.add_row().cells
        left.width, right.width = Inches(1.8), Inches(4.3)
        left.text = label
        left.paragraphs[0].runs[0].bold = True
        right.text = ""
        if value.startswith("https://"):
            hyperlink(right.paragraphs[0], value, value)
        else:
            right.paragraphs[0].add_run(value)
    document.add_paragraph()
    note = document.add_paragraph("The repository contains the implementation, experiment settings, saved predictions, evaluation results, and README.")
    note.paragraph_format.space_before = Pt(6)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()

